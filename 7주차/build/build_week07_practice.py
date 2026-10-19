#!/usr/bin/env python3
"""7주차 연습문제 세트 발표자료 빌드 (1부 문제 → 2부 풀이).

사용법: python3 build_week07_practice.py [작업폴더]
 - 수식: LaTeX(pdflatex, standalone) → 600dpi 투명 PNG
 - 도식: 편집 가능한 PPT 도형
 - 산출: 전자기학_7주차_연습문제_풀이_발표.pptx
디자인 토큰과 용어는 6.강의안/지침/PPT_PDF_제작_지침.md를 따른다.
"""
import math, re, subprocess, sys, hashlib
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
from PIL import Image

WORK = Path(sys.argv[1] if len(sys.argv) > 1 else ".").resolve()
EQ = WORK / "eq"; EQ.mkdir(exist_ok=True)
OUT = WORK / "전자기학_7주차_연습문제_풀이_발표.pptx"

C = dict(black="050505", ink="191919", white="FFFFFF", paper="F7F8FA", navy="0B1F3A",
         blue="2563EB", bluesoft="E0F0FE", gray="667085", graydark="344054", hair="D0D5DD")
KO = "Apple SD Gothic Neo"
MATH = "Times New Roman"
FOOT = "전자기학 2026 · WEEK 07 · 연습"
SRC_RANGE = "사용자 지정 7주차 범위: §3.1 복습, §3.2, §3.3 · 7주차_연습문제_세트.pdf (10월 19일 연습)"

def px(v): return Emu(int(round(v * 9525)))
def rgb(h): return RGBColor.from_string(h)

prs = Presentation()
prs.slide_width, prs.slide_height = px(1280), px(720)
BLANK = prs.slide_layouts[6]

# ---------------------------------------------------------------- equations
PRE = r"""\documentclass[border=3pt]{standalone}
\usepackage{amsmath,amssymb,bm,xcolor}
\newcommand{\E}{\mathbf E}\newcommand{\rhat}{\hat{\mathbf r}}\newcommand{\zhat}{\hat{\mathbf z}}
\newcommand{\xhat}{\hat{\mathbf x}}\newcommand{\that}{\hat{\boldsymbol\theta}}
\begin{document}\color[HTML]{%s}$\displaystyle %s$\end{document}"""
EQSRC = {}

def eq_png(key, latex, color="ink"):
    name = f"{key}_{color}"
    EQSRC[key] = latex
    h = hashlib.md5((latex + color).encode()).hexdigest()[:8]
    png = EQ / f"{name}_{h}.png"
    if not png.exists():
        tex = EQ / f"{name}.tex"
        tex.write_text(PRE % (C[color], latex))
        r = subprocess.run(["pdflatex", "-interaction=nonstopmode", "-halt-on-error", tex.name],
                           cwd=EQ, capture_output=True, text=True)
        if r.returncode:
            raise RuntimeError(f"LaTeX 오류 {key}:\n" + r.stdout[-1500:])
        subprocess.run(["pdftocairo", "-png", "-transp", "-r", "600", "-singlefile",
                        f"{name}.pdf", f"{name}_{h}"], cwd=EQ, check=True)
    return png

LAST = {}
SCALE = 2.45  # LaTeX 10pt → 화면 약 24.5pt

def equation(slide, key, latex, left, top, width, height=None, color="ink", scale=SCALE,
             align="left", valign="top"):
    png = eq_png(key, latex, color)
    w, h = Image.open(png).size
    wpx, hpx = w / 600 * 96 * scale, h / 600 * 96 * scale
    if wpx > width:
        f = width / wpx; wpx, hpx = wpx * f, hpx * f
    if height and hpx > height:
        f = height / hpx; wpx, hpx = wpx * f, hpx * f
    x = left + {"left": 0, "center": (width - wpx) / 2, "right": width - wpx}[align]
    y = top
    if height and valign == "middle": y = top + (height - hpx) / 2
    pic = slide.shapes.add_picture(str(png), px(x), px(y), px(wpx), px(hpx))
    pic.name = f"eq:{key}"
    # 대체 텍스트
    pic._element.nvPicPr.cNvPr.set("descr", latex)
    slide._eqs.append(latex)
    LAST["w"] = wpx
    return hpx

# ---------------------------------------------------------------- text
TOKEN = re.compile(r"(\$[^$]+\$)")
SUBSUP = re.compile(r"([_^])\{([^}]*)\}|([_^])(.)")

def _run(p, text, size, bold, color, font, italic=False, baseline=None):
    r = p.add_run(); r.text = text
    f = r.font; f.size = Pt(size); f.bold = bold; f.italic = italic; f.color.rgb = rgb(C[color])
    f.name = font
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", font)
    if baseline: rPr.set("baseline", str(baseline))
    return r

FUNC = re.compile(r"(sinh|cosh|sin|cos)")
def _math_runs(p, s, size, bold, color):
    parts = FUNC.split(s)
    if len(parts) > 1:
        for i, part in enumerate(parts):
            if not part: continue
            if i % 2: _run(p, part, size, bold, color, MATH, italic=False)
            else: _math_runs(p, part, size, bold, color)
        return
    pos = 0
    for m in SUBSUP.finditer(s):
        if m.start() > pos:
            _run(p, s[pos:m.start()], size, bold, color, MATH, italic=True)
        kind = m.group(1) or m.group(3); body = m.group(2) if m.group(1) else m.group(4)
        _run(p, body, size, bold, color, MATH, italic=True, baseline=-25000 if kind == "_" else 30000)
        pos = m.end()
    if pos < len(s):
        _run(p, s[pos:], size, bold, color, MATH, italic=True)

def text(slide, value, left, top, width, height, size=22, bold=False, color="ink",
         align="left", valign="top", spacing=1.12, fill=None, line=None, font=KO, para_gap=0):
    tb = slide.shapes.add_textbox(px(left), px(top), px(width), px(height))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[valign]
    bodyPr = tf._txBody.find(qn("a:bodyPr"))
    for ch in list(bodyPr): bodyPr.remove(ch)
    etree.SubElement(bodyPr, qn("a:noAutofit"))
    if fill:
        tb.fill.solid(); tb.fill.fore_color.rgb = rgb(C[fill])
    for i, para in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
        p.line_spacing = spacing
        if para_gap and i: p.space_before = Pt(para_gap)
        for tok in TOKEN.split(para):
            if not tok: continue
            if tok.startswith("$"):
                _math_runs(p, tok[1:-1], size, bold, color)
            else:
                _run(p, tok, size, bold, color, font)
    return tb

def est_height(value, size, width, spacing=1.12, para_gap=0):
    lines = 0
    em = size / 0.75
    for para in value.split("\n"):
        plain = re.sub(r"[$]", "", re.sub(r"[_^]\{([^}]*)\}", r"\1", para))
        w = 0
        for ch in plain:
            if ord(ch) > 0x3000: w += em * 0.95
            elif ch == " ": w += em * 0.28
            else: w += em * 0.52
        lines += max(1, math.ceil(w / width))
    return lines * em * 1.2 * spacing + para_gap / 0.75 * (value.count("\n"))

# ---------------------------------------------------------------- shapes
def rect(slide, l, t, w, h, fill=None, line=None, lw=0, shape=MSO_SHAPE.RECTANGLE, dash=False):
    s = slide.shapes.add_shape(shape, px(l), px(t), px(w), px(h))
    if fill: s.fill.solid(); s.fill.fore_color.rgb = rgb(C[fill])
    else: s.fill.background()
    if line:
        s.line.color.rgb = rgb(C[line]); s.line.width = Pt(lw)
        if dash:
            from pptx.enum.dml import MSO_LINE
            s.line.dash_style = MSO_LINE.DASH
    else: s.line.fill.background()
    s.shadow.inherit = False
    st = s._element.find(qn("p:style"))
    if st is not None: s._element.remove(st)
    return s

def hline(slide, x, y, w, color="hair", weight=1):
    return rect(slide, x, y - weight / 2, w, weight, fill=color)

def seg(slide, x1, y1, x2, y2, color="ink", lw=2.0, head=False, tail=False, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(x1), px(y1), px(x2), px(y2))
    c.line.color.rgb = rgb(C[color]); c.line.width = Pt(lw)
    st = c._element.find(qn("p:style"))
    if st is not None: c._element.remove(st)
    ln = c.line._get_or_add_ln()
    if dash:
        pd = etree.SubElement(ln, qn("a:prstDash")); pd.set("val", "dash")
    if tail:
        e = etree.SubElement(ln, qn("a:headEnd")); e.set("type", "triangle"); e.set("w", "lg"); e.set("len", "lg")
    if head:
        e = etree.SubElement(ln, qn("a:tailEnd")); e.set("type", "triangle"); e.set("w", "lg"); e.set("len", "lg")
    return c

def circle(slide, cx, cy, r, fill=None, line="ink", lw=2, dash=False):
    return rect(slide, cx - r, cy - r, 2 * r, 2 * r, fill, line, lw, MSO_SHAPE.OVAL, dash)

def charge(slide, cx, cy, label, positive=True, image=False, r=27):
    if image:
        circle(slide, cx, cy, r, fill="white", line="gray", lw=2.25, dash=True)
        col = "gray"
    else:
        circle(slide, cx, cy, r, fill="blue" if positive else "graydark", line=None)
        col = "white"
    text(slide, f"${label}$", cx - 60, cy - 17, 120, 34, size=19, bold=True, color=col, align="center", valign="middle")

def mlabel(slide, s, x, y, w=120, h=34, size=22, color="ink", align="left", bold=False):
    return text(slide, s, x, y, w, h, size=size, color=color, align=align, valign="middle", bold=bold)

# ---------------------------------------------------------------- slide frame
def new_slide(dark=False, bg=None):
    s = prs.slides.add_slide(BLANK)
    s._eqs = []
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(C[bg or ("navy" if dark else "white")])
    return s

PAGE = [0]
def footer(s, dark=False):
    PAGE[0] += 1
    hline(s, 64, 672, 1152, "graydark" if dark else "hair", 1)
    text(s, FOOT, 64, 683, 500, 22, size=12, bold=True, color="gray" if dark else "graydark")
    text(s, f"{PAGE[0]:02d}", 1116, 683, 100, 22, size=12, bold=True, color="gray" if dark else "graydark", align="right")

def header(s, kicker, title, dark=False, tsize=32):
    text(s, kicker, 64, 36, 1152, 24, size=14, bold=True, color="bluesoft" if dark else "blue")
    text(s, title, 64, 66, 1152, 56, size=tsize, bold=True, color="white" if dark else "ink", valign="middle")
    footer(s, dark)

def notes(s, teaching, sections):
    eqs = "\n".join(f"- {e}" for e in s._eqs)
    body = teaching
    if eqs: body += "\n\n[Equation source]\n" + eqs
    body += (f"\n\n[Sources]\n- Griffiths, Introduction to Electrodynamics, 4e, {sections}\n"
             f"- {SRC_RANGE}\n- 전자기학 강의안 PPT·PDF 제작 지침 v1.8\n[/Sources]")
    s.notes_slide.notes_text_frame.text = body

def result_box(s, key, latex, top=566, height=92, label="결과", scale=SCALE, left=64, width=1152, tail=None, tail_w=0):
    rect(s, left, top, width, height, fill="bluesoft")
    rect(s, left, top, 8, height, fill="blue")
    text(s, label, left + 28, top, 110, height, size=22, bold=True, color="blue", valign="middle")
    equation(s, key, latex, left + 140, top + 4, width - 160 - tail_w, height - 8, scale=scale, valign="middle")
    if tail:
        x = left + 140 + LAST["w"] + 16
        text(s, tail, x, top, left + width - x - 12, height, size=21, bold=True, color="navy", valign="middle")

def stack(s, items, left, top, width, gap=14):
    """items: ('t', text, size) | ('e', key, latex[, scale, align]) | ('g', px)"""
    y = top
    for it in items:
        if it[0] == "t":
            size = it[2] if len(it) > 2 else 22
            bold = it[3] if len(it) > 3 else False
            h = est_height(it[1], size, width)
            text(s, it[1], left, y, width, h + 4, size=size, bold=bold)
            y += h + gap
        elif it[0] == "e":
            sc = it[3] if len(it) > 3 else SCALE
            al = it[4] if len(it) > 4 else "left"
            h = equation(s, it[1], it[2], left, y, width, scale=sc, align=al)
            y += h + gap
        elif it[0] == "g":
            y += it[1]
    return y

def step_label(s, n, x, y):
    circle(s, x + 14, y + 16, 14, fill="navy", line=None)
    text(s, str(n), x, y + 2, 28, 28, size=15, bold=True, color="white", align="center", valign="middle")

# ---------------------------------------------------------------- diagrams
def plane_diagram(s, x0=110, ground_y=470, qy=280, image=False, rho=False, shade=False, width=440):
    cx = x0 + width / 2
    if shade:
        rect(s, x0, ground_y + 6, width, 150, fill="paper")
        text(s, "도체 쪽: 장 없음", cx + 40, ground_y + 112, 180, 28, size=17, color="gray")
    # z축
    seg(s, cx, ground_y + (165 if image else 40), cx, qy - 85, color="gray", lw=1.5, head=True)
    mlabel(s, "$z$", cx + 12, qy - 105, 40, 34, size=24)
    # 접지 평면
    rect(s, x0, ground_y - 3, width, 6, fill="graydark")
    for i in range(12):
        x = x0 + 18 + i * (width - 36) / 11
        seg(s, x, ground_y + 3, x - 14, ground_y + 20, color="graydark", lw=1.25)
    mlabel(s, "$z=0$, $V=0$", x0 + width - 150, ground_y + 26, 150, 34, size=20, align="right")
    charge(s, cx, qy, "+q", True)
    # 높이 a 치수선
    dx = cx - 70
    seg(s, dx, ground_y - 6, dx, qy, color="ink", lw=1.25, head=True, tail=True)
    mlabel(s, "$a$", dx - 46, (ground_y + qy) / 2 - 17, 36, 34, size=24, align="right")
    if image:
        iy = 2 * ground_y - qy
        charge(s, cx, iy, "−q", image=True)
        seg(s, dx, ground_y + 6, dx, iy, color="gray", lw=1.25, head=True, tail=True)
        mlabel(s, "$a$", dx - 46, (ground_y + iy) / 2 - 17, 36, 34, size=24, align="right", color="gray")
    if rho:
        seg(s, cx, ground_y - 22, cx + 150, ground_y - 22, color="blue", lw=2, head=True)
        mlabel(s, "$ρ$", cx + 155, ground_y - 42, 40, 34, size=24, color="blue")

def sphere_diagram(s, cx=330, cy=380, R=150, ax=620, image=False, center=None, ground=True, qlabel="+q"):
    circle(s, cx, cy, R, fill="paper", line="graydark", lw=2.5)
    seg(s, cx - R - 40, cy, ax + 60, cy, color="gray", lw=1.5, head=True)
    mlabel(s, "$z$", ax + 46, cy - 44, 36, 34, size=24)
    circle(s, cx, cy, 4, fill="ink", line=None)
    # 반지름 R (위쪽 대각선)
    ang = math.radians(125)
    ex, ey = cx + R * math.cos(ang), cy - R * math.sin(ang)
    seg(s, cx, cy, ex, ey, color="ink", lw=1.5, head=True)
    mlabel(s, "$R$", (cx + ex) / 2 + 6, (cy + ey) / 2 - 26, 40, 34, size=24)
    charge(s, ax, cy, qlabel, True)
    # a 치수선 (아래)
    dy = cy + R + 34
    seg(s, cx, cy + 10, cx, dy + 8, color="hair", lw=1)
    seg(s, ax, cy + 30, ax, dy + 8, color="hair", lw=1)
    seg(s, cx, dy, ax, dy, color="ink", lw=1.25, head=True, tail=True)
    mlabel(s, "$a$", (cx + ax) / 2 - 18, dy + 4, 36, 34, size=24, align="center")
    if image:
        bx = cx + R * R / (ax - cx)
        charge(s, bx, cy, "q′", image=True, r=24)
        seg(s, cx, cy - 40, bx, cy - 40, color="gray", lw=1.25, head=True, tail=True)
        mlabel(s, "$b$", (cx + bx) / 2 - 18, cy - 78, 36, 34, size=24, align="center", color="gray")
    if center:
        charge(s, cx, cy, center, image=True, r=24)
    if ground:
        gx, gy = cx - R * 0.70, cy + R * 0.71
        seg(s, gx, gy, gx - 40, gy + 40, color="graydark", lw=1.5)
        for i, w in enumerate((36, 24, 12)):
            seg(s, gx - 40 - w / 2 - 14, gy + 40 + i * 8, gx - 40 + w / 2 - 14, gy + 40 + i * 8, color="graydark", lw=1.5)
        text(s, "접지", gx - 84, gy + 64, 60, 28, size=17, color="graydark", align="center")

def rect_diagram(s, x0=150, y0=200, w=380, h=300, right="V(b,y)=f(y)", rcolor="blue"):
    rect(s, x0, y0, w, h, fill="white", line=None)
    seg(s, x0, y0 + h, x0 + w + 60, y0 + h, color="gray", lw=1.5, head=True)
    seg(s, x0, y0 + h, x0, y0 - 50, color="gray", lw=1.5, head=True)
    mlabel(s, "$x$", x0 + w + 62, y0 + h - 18, 30, 34, size=24)
    mlabel(s, "$y$", x0 - 10, y0 - 92, 30, 34, size=24)
    for (a, b, c, d) in ((x0, y0, x0 + w, y0), (x0, y0 + h, x0 + w, y0 + h), (x0, y0, x0, y0 + h)):
        seg(s, a, b, c, d, color="graydark", lw=3)
    seg(s, x0 + w, y0, x0 + w, y0 + h, color=rcolor, lw=4.5)
    mlabel(s, "$V=0$", x0 + w / 2 - 60, y0 - 46, 120, 34, size=21, align="center")
    mlabel(s, "$V=0$", x0 + w / 2 - 60, y0 + h + 10, 120, 34, size=21, align="center")
    mlabel(s, "$V=0$", x0 - 100, y0 + h / 2 - 17, 86, 34, size=21, align="right")
    mlabel(s, f"${right}$", x0 + w + 14, y0 + h / 2 - 17, 210, 34, size=21, color=rcolor)
    mlabel(s, "$0$", x0 - 30, y0 + h + 6, 24, 30, size=19, color="gray")
    mlabel(s, "$b$", x0 + w - 12, y0 + h + 40, 24, 30, size=21)
    mlabel(s, "$a$", x0 - 32, y0 - 16, 24, 30, size=21)

def cos_sphere(s, cx=320, cy=390, R=165, inside_field=False):
    circle(s, cx, cy, R, fill="paper", line="blue", lw=3)
    seg(s, cx, cy + R + 40, cx, cy - R - 70, color="gray", lw=1.5, head=True)
    mlabel(s, "$z$", cx + 10, cy - R - 92, 40, 34, size=24)
    ang = math.radians(55)
    ex, ey = cx + R * math.cos(ang), cy - R * math.sin(ang)
    seg(s, cx, cy, ex, ey, color="ink", lw=1.5, head=True)
    mlabel(s, "$θ$", cx + 8, cy - 92, 36, 34, size=22)
    mlabel(s, "$R$", (cx + ex) / 2 + 8, (cy + ey) / 2 - 8, 30, 34, size=22)
    mlabel(s, "$V(R,θ)=V_{0} cos θ$", cx + R * 0.74, cy + R * 0.80, 240, 34, size=21, color="blue")
    if inside_field:
        for dx in (-80, 0, 80):
            seg(s, cx + dx, cy - 70, cx + dx, cy + 80, color="blue", lw=2.5, head=True)
        text(s, "내부: 균일한 전기장", cx - R - 20, cy + R + 14, 220, 28, size=17, bold=True, color="blue")

# ---------------------------------------------------------------- problem slide helper
def problem_slide(code, pts, section, title, intro, parts, diagram=None, sections="§3.2", teach="",
                  eqs=(), text_w=None, size=21):
    s = new_slide()
    header(s, f"{section} · {code} · {pts}점", title)
    tw = text_w or (600 if diagram else 1152)
    tx = 616 if diagram else 64
    y = 148
    h = est_height(intro, size, tw)
    text(s, intro, tx, y, tw, h + 6, size=size)
    y += h + 14
    for k, l in eqs:
        y += equation(s, k, l, tx, y, tw, align="center") + 12
    for part in parts:
        lab, body = part[0], part[1]
        h = est_height(body, size, tw - 48)
        text(s, lab, tx, y, 44, 32, size=size, bold=True, color="blue")
        text(s, body, tx + 48, y, tw - 48, h + 6, size=size)
        y += h + 12
        if len(part) > 2:
            y += equation(s, part[2][0], part[2][1], tx + 48, y, tw - 48, align="center") + 10
    if y > 660: print(f"[경고] {code} 문제 장표 넘침 y={y:.0f}")
    if diagram: diagram(s)
    notes(s, teach or f"{code} 출제 화면. 학생이 먼저 실제 영역과 경계조건을 적도록 안내한다.", sections)
    return s

def sol_slide(kicker, title, sections, teach):
    s = new_slide(); header(s, kicker, title)
    s._teach = teach; s._sec = sections
    return s

def done(s):
    notes(s, s._teach, s._sec)

def section_slide(kicker, title, sub, teach):
    s = new_slide(dark=True)
    text(s, kicker, 64, 36, 1152, 24, size=14, bold=True, color="bluesoft")
    text(s, title, 64, 230, 1152, 90, size=48, bold=True, color="white", align="center", valign="middle")
    hline(s, 490, 345, 300, "blue", 5)
    text(s, sub, 140, 380, 1000, 90, size=26, bold=True, color="bluesoft", align="center", valign="top")
    footer(s, True)
    notes(s, teach, "§3.1–§3.3")
    return s

K = r"\frac{1}{4\pi\epsilon_0}"

# =============================================================== 01 cover
s = new_slide(bg="black")
rect(s, 0, 0, 18, 720, fill="blue")
text(s, "ELECTROMAGNETISM 2026 · WEEK 07 · PRACTICE", 80, 64, 900, 30, size=16, bold=True, color="white")
text(s, "7주차 연습문제\n영상법과 변수 분리법", 80, 150, 820, 200, size=54, bold=True, color="white", spacing=1.0)
hline(s, 82, 395, 610, "blue", 5)
text(s, "Griffiths §3.1 복습 · §3.2 · §3.3", 80, 425, 800, 44, size=26, bold=True, color="bluesoft")
text(s, "10월 19일 연습 · 1부 문제 → 2부 풀이", 80, 560, 800, 36, size=22, color="white")
# 표지 도식: 접지 평면 + 실제/영상 전하
rect(s, 930, 397, 290, 6, fill="white")
circle(s, 1075, 260, 30, fill="blue", line=None)
text(s, "$+q$", 1015, 243, 120, 34, size=20, bold=True, color="white", align="center", valign="middle")
for x in (985, 1030, 1075, 1120, 1165):
    dx = x - 1075; n = math.hypot(dx, 137)
    seg(s, 1075 + dx / n * 32, 260 + 137 / n * 32, x, 390, color="blue", lw=2.5, head=True)
circle(s, 1075, 540, 30, fill="black", line="gray", lw=2.25, dash=True)
text(s, "$−q$", 1015, 523, 120, 34, size=20, bold=True, color="gray", align="center", valign="middle")
text(s, "$V=0$", 1120, 410, 100, 30, size=19, color="bluesoft", align="right")
s._eqs = []
notes(s, "표지. 7주차 연습문제 세트(P1–P8)를 1부에서 출제하고 2부에서 풀이한다. 도식은 접지 평면과 영상전하를 예고한다.", "§3.1–§3.3")
PAGE[0] += 1

# =============================================================== 02 guide
s = new_slide(); header(s, "W07 / PRACTICE GUIDE", "연습 진행 방법과 풀이 원칙")
text(s, "풀이 원칙", 64, 150, 520, 34, size=24, bold=True, color="blue")
items = ["실제 영역과 경계조건을 먼저 적는다.",
         "영상전하는 반드시 실제 영역 밖에 둔다.",
         "변수 분리법에서는 각 항이 라플라스 방정식과 모든 경계조건을 만족하는지 확인한다.",
         "부호·차원·극한으로 최종 답을 검산한다."]
y = 200
for i, it in enumerate(items):
    step_label(s, i + 1, 64, y)
    h = est_height(it, 22, 480)
    text(s, it, 108, y, 480, h + 4, size=22)
    y += max(h, 34) + 22
text(s, "문제 구성 (100점)", 660, 150, 556, 34, size=24, bold=True, color="blue")
rows = [("A", "접지된 무한 평면의 영상법", "P1–P3", "34점"),
        ("B", "구면 경계의 영상법", "P4–P5", "25점"),
        ("C", "직교좌표계의 변수 분리법", "P6–P7", "27점"),
        ("D", "구 좌표계의 변수 분리법", "P8", "14점")]
y = 205
for a, b, c, d in rows:
    hline(s, 660, y - 8, 556, "hair", 1)
    text(s, a, 660, y, 40, 40, size=24, bold=True, color="navy")
    text(s, b, 706, y, 330, 40, size=21, bold=True)
    text(s, c, 1040, y, 90, 40, size=20, color="graydark")
    text(s, d, 1130, y, 86, 40, size=20, bold=True, color="blue", align="right")
    y += 62
hline(s, 660, y - 8, 556, "hair", 1)
rect(s, 660, 488, 556, 128, fill="paper")
text(s, "1부 (3–11장)  문제만 제시 → 각자 풀이\n2부 (12장부터)  문제별 풀이·검산 공개", 684, 506, 520, 92, size=21, bold=True, spacing=1.35)
notes(s, "연습 시간 운영: 1부의 문제 화면을 띄워 두고 학생이 먼저 풀게 한 뒤 2부 풀이로 넘어간다. 진공의 유전율은 ε₀.", "§3.1–§3.3")

# =============================================================== 03 section
section_slide("PART 1 / PROBLEMS", "1부 · 문제", "먼저 스스로 풀어 봅시다. 풀이는 2부에서 공개합니다.",
              "문제 출제 구간. 각 문제 화면에는 완전한 문제 문장을 싣는다.")

# =============================================================== 04–11 problems
problem_slide("P1", 12, "A. 접지된 무한 평면의 영상법", "점전하와 접지된 평면의 전위",
    "접지된 무한 도체평면이 $z=0$에 있고, 점전하 $+q$가 $(0, 0, a)$에 놓여 있다. 실제 영역은 $z>0$이다.",
    [("(a)", "영상전하의 위치와 크기를 쓰고 그 선택 이유를 설명하시오."),
     ("(b)", "원통좌표계 $(ρ, φ, z)$로 실제 영역의 전위 $V(ρ, z)$를 구하시오."),
     ("(c)", "$V(ρ, 0)=0$과 $V→0$ $(r→∞)$를 확인하고 유일성 정리로 결과를 정당화하시오.")],
    diagram=lambda s: plane_diagram(s, x0=90, width=440, rho=True), sections="§3.1, §3.2.1", size=23)

problem_slide("P2", 12, "A. 접지된 무한 평면의 영상법", "평면의 유도 면전하",
    "P1의 전위를 사용한다.",
    [("(a)", "다음 식으로 면전하밀도를 계산하시오.", ("p2sig", r"\sigma(\rho)=-\epsilon_0\left.\frac{\partial V}{\partial z}\right|_{z=0^+}")),
     ("(b)", "$|σ|$가 가장 큰 위치와 $ρ≫a$에서의 감소율을 구하시오."),
     ("(c)", "다음 적분으로 유도된 총전하를 구하고 전하 보존의 관점에서 해석하시오.", ("p2Q", r"Q_{\rm ind}=\int_0^\infty\sigma(\rho)\,2\pi\rho\,\mathrm d\rho"))],
    sections="§3.2.1", text_w=1152, size=23)

problem_slide("P3", 10, "A. 접지된 무한 평면의 영상법", "실제 전하가 받는 힘과 일",
    "P1의 점전하에 대하여 다음을 구하시오.",
    [("(a)", "실제 전하가 받는 힘을 영상전하가 만드는 장으로부터 구하시오."),
     ("(b)", "전하를 무한대에서 높이 $a$까지 천천히 가져올 때 외력이 한 일을 구하시오."),
     ("(c)", "두 점전하의 상호작용 에너지를 그대로 사용하면 에너지가 두 배가 되는 이유를 설명하시오.")],
    diagram=lambda s: plane_diagram(s, x0=90, width=440), sections="§3.2.1, §3.2.3", size=23)

problem_slide("P4", 15, "B. 구면 경계의 영상법", "점전하와 접지된 도체구",
    "반지름 $R$인 접지된 도체구의 중심을 원점에 두고, 점전하 $+q$를 $z$축 위 $z=a$에 놓는다. 단, $a>R$이다.",
    [("(a)", "구 안에 놓을 영상전하의 위치 $b$와 크기 $q′$를 구하시오."),
     ("(b)", "$r≥R$에서의 전위를 쓰고 $r=R$에서 모든 $θ$에 대해 $V=0$임을 보이시오."),
     ("(c)", "구에 유도된 총전하를 구하시오."),
     ("(d)", "실제 전하가 받는 힘의 크기와 방향을 구하시오.")],
    diagram=lambda s: sphere_diagram(s, cx=250, cy=390, R=135, ax=490), sections="§3.2.2", size=22)

problem_slide("P5", 10, "B. 구면 경계의 영상법", "중성인 고립 도체구",
    "P4의 접지선을 제거하고 도체구의 총전하를 0으로 유지한다.",
    [("(a)", "P4의 영상전하에 더하여 구 중심에 놓아야 할 영상전하 $q_{0}′$를 구하시오."),
     ("(b)", "구 바깥의 전위를 쓰고 구면이 등전위임을 확인하시오."),
     ("(c)", "접지된 경우와 비교하여 무한대에서 전위의 가장 낮은 차수 항이 어떻게 달라지는지 설명하시오.")],
    diagram=lambda s: (sphere_diagram(s, cx=250, cy=390, R=135, ax=490, ground=False),
                       text(s, "총전하 0 (고립)", 150, 210, 200, 30, size=18, bold=True, color="graydark", align="center")),
    sections="§3.2.2", size=23)

problem_slide("P6", 13, "C. 직교좌표계의 변수 분리법", "하나의 사인 모드로 주어진 경계",
    "전하가 없는 직사각형 영역 $0<x<b$, $0<y<a$에서 다음 경계조건이 주어진다.",
    [("(a)", "$V=XY$를 대입하여 분리된 두 상미분방정식을 구하시오."),
     ("(b)", "네 경계조건을 만족하는 전위 $V(x, y)$를 구하시오."),
     ("(c)", "$x=b/2$, $y=a/2$에서 전위와 전기장을 구하시오.")],
    eqs=[("p6bc", r"V(0,y)=V(x,0)=V(x,a)=0,\qquad V(b,y)=V_0\sin\frac{\pi y}{a}")],
    diagram=lambda s: rect_diagram(s, x0=130, y0=250, w=280, h=260, right="V_{0} sin(πy/a)"),
    sections="§3.3.1", size=22)

problem_slide("P7", 14, "C. 직교좌표계의 변수 분리법", "일정한 경계전위의 푸리에 전개",
    "P6과 같은 영역에서 세 경계는 0이고 $V(b, y)=V_{0}$ $(0<y<a)$이다.",
    [("(a)", "$V_{0}$를 사인 급수로 전개하여 짝수 모드가 사라짐을 보이시오."),
     ("(b)", "영역 안의 전위를 무한급수로 나타내시오."),
     ("(c)", "급수의 각 항이 라플라스 방정식과 세 영 경계를 만족함을 확인하시오."),
     ("(d)", "모서리에서 경계값이 불연속일 때 급수가 어떤 값으로 수렴하는지 설명하시오.")],
    diagram=lambda s: rect_diagram(s, x0=130, y0=250, w=280, h=260, right="V_{0}"),
    sections="§3.3.1, Example 3.3", size=22)

problem_slide("P8", 14, "D. 구 좌표계의 변수 분리법", "cos θ 경계가 고르는 모드",
    "반지름 $R$인 구면의 경계전위가 $V(R, θ)=V_{0} cos θ$로 주어진다.",
    [("(a)", "구 안쪽 $r<R$에서 원점에 유한한 전위를 구하고 전기장을 해석하시오."),
     ("(b)", "구 바깥 $r>R$에서 $V→0$ $(r→∞)$인 전위를 구하시오."),
     ("(c)", "두 해에서 필요한 르장드르 다항식의 차수 $ℓ$을 밝히시오."),
     ("(d)", "바깥 해의 먼 거리 거동으로부터 등가 쌍극자 모멘트를 구하시오.")],
    diagram=lambda s: cos_sphere(s, cx=290, cy=400, R=150), sections="§3.3.2, Example 3.6–3.7", size=22)

# =============================================================== 12 section
section_slide("PART 2 / SOLUTIONS", "2부 · 풀이", "실제 영역 → 경계조건 → 풀이 → 검산 순서로 확인합니다.",
              "풀이 공개 구간. 각 장표는 상단 핵심식, 중앙 단계, 하단 결과 순서로 읽는다.")

# =============================================================== P1
s = sol_slide("풀이 P1 (a) · 영상전하", "거울 위치의 반대 전하가 평면 전위를 0으로 만든다", "§3.2.1",
              "영상전하는 실제 영역 밖(z<0)에 있으므로 z>0의 전하분포를 바꾸지 않는다. 평면의 모든 점은 두 전하에서 같은 거리에 있다.")
plane_diagram(s, x0=90, width=440, ground_y=380, qy=230, image=True)
stack(s, [("t", "실제 영역 $z>0$ · 경계조건 $V(z=0)=0$, $V(∞)=0$", 22, True),
          ("t", "평면 위의 점은 $(0,0,a)$와 $(0,0,−a)$에서 거리가 같다:", 22),
          ("e", "p1d", r"\sqrt{\rho^2+a^2}\;=\;\sqrt{\rho^2+(-a)^2}"),
          ("t", "크기가 같고 부호가 반대인 전하를 두면 두 전위가 평면 전체에서 상쇄된다. 영상전하는 $z<0$에 있어 실제 영역의 전하분포를 바꾸지 않는다.", 22)],
      616, 150, 600)
result_box(s, "p1res", r"q'=-q,\qquad \mathbf r_{q'}=(0,0,-a)", left=616, width=600)
done(s)

s = sol_slide("풀이 P1 (b)–(c) · 전위와 정당화", "실제 영역의 전위는 두 점전하 전위의 합이다", "§3.1, §3.2.1",
              "z<0(도체 쪽)의 전위는 0이다. 무한대 거동은 쌍극자형 1/r^2 감소. 유일성 정리로 후보해를 실제 해로 확정한다.")
y = stack(s, [("e", "p1V", r"V(\rho,z)=\frac{q}{4\pi\epsilon_0}\left[\frac{1}{\sqrt{\rho^2+(z-a)^2}}-\frac{1}{\sqrt{\rho^2+(z+a)^2}}\right]\quad(z\ge 0)", SCALE, "center")], 64, 145, 1152)
hline(s, 64, y + 4, 1152, "hair", 1)
step_label(s, 1, 64, y + 22)
stack(s, [("t", "$z=0$: 두 분모가 모두 $√(ρ^{2}+a^{2})$ → $V(ρ,0)=0$", 22)], 108, y + 22, 520)
step_label(s, 2, 64, y + 92)
stack(s, [("t", "$r→∞$: 두 항의 차가 남는다", 22),
          ("e", "p1far", r"V\simeq\frac{1}{4\pi\epsilon_0}\frac{2qa\,z}{r^3}\to 0")], 108, y + 92, 520)
step_label(s, 3, 660, y + 22)
stack(s, [("t", "제1 유일성 정리", 22, True),
          ("t", "$z>0$에서 점전하 $q$에 대한 푸아송 방정식을 만족하고 모든 경계조건($z=0$, 무한대)을 만족하므로 이 후보해가 유일한 실제 해이다.", 22)], 704, y + 22, 512)
result_box(s, "p1res2", r"V(\rho,0)=0,\qquad V\to0\ \ (r\to\infty)", tail="→ 유일성 정리에 따라 $z>0$의 유일한 해", tail_w=420)
done(s)

# =============================================================== P2
s = sol_slide("풀이 P2 (a) · 유도 면전하밀도", "법선 미분이 면전하밀도를 결정한다", "§3.2.1",
              "σ = −ε0 ∂V/∂n. 표면 바로 위의 장 E_z(0+) = σ/ε0 와 부호가 일치하는지 확인한다(장선이 평면을 향해 아래로).")
y = 140
step_label(s, 1, 64, y + 10)
y = stack(s, [("t", "$σ=−ε_{0}(∂V/∂z)_{z=0+}$ 계산: 각 항을 $z$로 미분한다", 22),
              ("e", "p2d1", r"\frac{\partial}{\partial z}\frac{1}{\sqrt{\rho^2+(z\mp a)^2}}=-\frac{z\mp a}{[\rho^2+(z\mp a)^2]^{3/2}}")], 108, y + 10, 1100)
step_label(s, 2, 64, y + 6)
stack(s, [("t", "$z=0$을 대입하면 두 항이 같은 값을 준다", 22),
          ("e", "p2d2", r"\left.\frac{\partial V}{\partial z}\right|_{0^+}=\frac{q}{4\pi\epsilon_0}\left[\frac{a}{(\rho^2+a^2)^{3/2}}+\frac{a}{(\rho^2+a^2)^{3/2}}\right]")], 108, y + 6, 1100)
result_box(s, "p2res", r"\sigma(\rho)=-\frac{qa}{2\pi(\rho^2+a^2)^{3/2}}")
done(s)

s = sol_slide("풀이 P2 (b)–(c) · 분포와 총전하", "유도전하는 전하 바로 아래에 모이고 총량은 −q이다", "§3.2.1",
              "총 유도전하 −q는 접지선을 통해 땅에서 공급된다. +q에서 나간 선속이 z<0으로 새지 않고 모두 평면에서 끝난다.")
# 막대 도식
x0, gy = 90, 330
rect(s, x0, gy - 3, 470, 6, fill="graydark")
for i in range(-6, 7):
    x = x0 + 235 + i * 34
    hh = 150 / (1 + 0.16 * i * i) ** 1.5
    rect(s, x - 5, gy + 6, 10, hh, fill="blue")
mlabel(s, "$|σ(ρ)|$", x0 + 300, gy + 60, 120, 34, size=22, color="blue")
mlabel(s, "$ρ=0$", x0 + 195, gy - 46, 80, 34, size=20, align="center")
text(s, "막대 길이 ∝ 면전하밀도의 크기", x0, gy + 175, 470, 30, size=17, color="gray", align="center")
stack(s, [("t", "(b) $|σ|$는 $ρ$에 대해 단조 감소 → 최대는 $ρ=0$", 22, True),
          ("e", "p2max", r"|\sigma|_{\max}=\frac{q}{2\pi a^2},\qquad \sigma\approx-\frac{qa}{2\pi\rho^3}\ \ (\rho\gg a)"),
          ("t", "(c) $u=ρ^{2}+a^{2}$로 치환", 22, True),
          ("e", "p2Qi", r"Q_{\rm ind}=-qa\int_0^\infty\frac{\rho\,\mathrm d\rho}{(\rho^2+a^2)^{3/2}}=-qa\cdot\frac1a=-q")],
      616, 150, 600)
result_box(s, "p2res2", r"|\sigma|_{\max}=\frac{q}{2\pi a^2}\ (\rho=0),\qquad \sigma\propto\rho^{-3},\qquad Q_{\rm ind}=-q", tail="(접지선으로 공급)", tail_w=220)
done(s)

# =============================================================== P3
s = sol_slide("풀이 P3 (a)–(b) · 힘과 일", "실제 전하는 영상전하 쪽으로 끌린다", "§3.2.1, §3.2.3",
              "z>0에서 유도전하가 만드는 장 = 영상전하의 장. 높이 z에서 외력은 전기력과 반대 방향으로 같은 크기.")
plane_diagram(s, x0=90, width=440, ground_y=380, qy=230, image=True)
seg(s, 360, 230, 360, 300, color="blue", lw=3, head=True)
text(s, "힘", 372, 268, 40, 30, size=19, bold=True, color="blue")
stack(s, [("t", "(a) 영상전하 $−q$가 거리 $2a$에서 만드는 장", 22, True),
          ("e", "p3F", r"\mathbf F=q\,\mathbf E_{\rm image}=-\frac{1}{4\pi\epsilon_0}\frac{q^2}{(2a)^2}\,\hat{\mathbf z}"),
          ("t", "(b) 높이 $z$에서 외력은 전기력과 크기가 같고 방향은 $+z$", 22, True),
          ("e", "p3W", r"W=\int_\infty^a\frac{q^2}{16\pi\epsilon_0 z^2}\,\mathrm dz=-\frac{q^2}{16\pi\epsilon_0 a}")],
      616, 150, 600)
result_box(s, "p3res", r"\mathbf F=-\frac{q^2}{16\pi\epsilon_0a^2}\hat{\mathbf z},\qquad W=-\frac{q^2}{16\pi\epsilon_0a}", left=616, width=600)
done(s)

s = sol_slide("풀이 P3 (c) · 에너지 계수 1/2", "장이 반공간에만 있으므로 에너지는 절반이다", "§3.2.3",
              "두 점전하계 에너지는 전 공간에 장이 있을 때의 값. 실제로는 z<0에 장이 없어 (ε0/2)∫E^2 가 절반. 또한 영상전하는 실제 전하를 따라 저절로 움직이므로 일을 받지 않는다.")
plane_diagram(s, x0=90, width=440, ground_y=375, qy=245, image=True, shade=True)
stack(s, [("e", "p3U2", r"W_{qq'}=-\frac{1}{4\pi\epsilon_0}\frac{q^2}{2a}=-\frac{q^2}{8\pi\epsilon_0 a}"),
          ("g", 4),
          ("t", "① 두 점전하계의 장은 전 공간에 있지만 실제 장은 $z>0$에만 있다. 장은 평면에 대해 대칭이므로 $(ε_{0}/2)∫E^{2}dτ$가 절반만 남는다.", 22),
          ("t", "② 실제 전하를 옮기면 영상전하는 저절로 따라 움직인다. 외력의 일은 실제 전하 하나에 대한 것뿐이다.", 22)],
      616, 150, 600)
result_box(s, "p3res2", r"W=\frac12\,W_{qq'}=-\frac{q^2}{16\pi\epsilon_0a}", left=616, width=600)
done(s)

# =============================================================== P4
s = sol_slide("풀이 P4 (a) · 영상전하의 위치와 크기", "구면의 두 점에서 V = 0을 요구한다", "§3.2.2",
              "θ=0과 θ=π 두 점의 조건으로 미지수 b, q′를 정한다. 이후 (b)에서 모든 θ에 대해 성립함을 확인한다.")
sphere_diagram(s, cx=230, cy=360, R=130, ax=480, image=True)
stack(s, [("t", "$θ=0$ (가까운 점)과 $θ=π$ (먼 점)", 22, True),
          ("e", "p4c", r"\frac{q}{a-R}+\frac{q'}{R-b}=0,\qquad \frac{q}{a+R}+\frac{q'}{R+b}=0"),
          ("t", "두 식을 연립하면", 22),
          ("e", "p4s", r"b=\frac{R^2}{a}<R,\qquad q'=-\frac{R}{a}\,q"),
          ("t", "$a>R$이므로 영상전하는 구 안, 즉 실제 영역 $r≥R$ 밖에 있다.", 22)],
      656, 150, 560)
result_box(s, "p4res", r"b=\frac{R^2}{a},\qquad q'=-\frac{R}{a}q", left=656, width=560)
done(s)

s = sol_slide("풀이 P4 (b) · 구 바깥의 전위", "거리비가 구면 전체에서 일정하다", "§3.2.2",
              "r=R에서 영상전하까지의 거리는 실제 전하까지 거리의 R/a배이므로 모든 θ에서 상쇄된다.")
y = stack(s, [("e", "p4V", r"V(r,\theta)=\frac{q}{4\pi\epsilon_0}\left[\frac{1}{\sqrt{r^2+a^2-2ar\cos\theta}}-\frac{R}{\sqrt{a^2r^2+R^4-2aR^2r\cos\theta}}\right]\ \ (r\ge R)", 2.3, "center")], 64, 145, 1152)
hline(s, 64, y + 4, 1152, "hair", 1)
step_label(s, 1, 64, y + 24)
y2 = stack(s, [("t", "$r=R$에서 둘째 항의 분모", 22),
               ("e", "p4b1", r"\sqrt{a^2R^2+R^4-2aR^3\cos\theta}=R\sqrt{a^2+R^2-2aR\cos\theta}")], 108, y + 24, 1100)
step_label(s, 2, 64, y2 + 6)
stack(s, [("t", "따라서 두 항이 같은 크기, 반대 부호가 되어 모든 $θ$에서 상쇄된다.", 22)], 108, y2 + 6, 1100)
result_box(s, "p4res2", r"V(R,\theta)=\frac{q}{4\pi\epsilon_0}\left[\frac{1}{\sqrt{R^2+a^2-2aR\cos\theta}}-\frac{1}{\sqrt{a^2+R^2-2aR\cos\theta}}\right]=0")
done(s)

s = sol_slide("풀이 P4 (c)–(d) · 유도전하와 힘", "유도 총전하는 영상전하와 같고 힘은 구를 향한다", "§3.2.2",
              "구를 감싸는 가우스 곡면(q 제외)의 선속 = q′/ε0. 접지되어 있으므로 유도 총전하가 −q일 필요는 없다. 검산: R→∞, a−R=d 고정이면 평면의 결과 q^2/[4πε0(2d)^2]로 돌아간다.")
sphere_diagram(s, cx=230, cy=360, R=130, ax=480, image=True)
seg(s, 480, 330, 420, 330, color="blue", lw=3, head=True)
text(s, "힘", 432, 292, 40, 30, size=19, bold=True, color="blue")
stack(s, [("t", "(c) 구 바깥의 장 = $q+q′$의 장 → 구를 감싼 가우스 곡면", 22, True),
          ("e", "p4Q", r"Q_{\rm ind}=q'=-\frac{R}{a}\,q"),
          ("t", "(d) 거리 $a-b=(a^{2}−R^{2})/a$", 22, True),
          ("e", "p4F", r"\mathbf F=\frac{1}{4\pi\epsilon_0}\frac{qq'}{(a-b)^2}\hat{\mathbf z}=-\frac{1}{4\pi\epsilon_0}\frac{q^2Ra}{(a^2-R^2)^2}\hat{\mathbf z}"),
          ],
      656, 150, 560)
result_box(s, "p4res3", r"|\mathbf F|=\frac{q^2Ra}{4\pi\epsilon_0(a^2-R^2)^2}", tail="구 중심 쪽을 향함 (인력)", tail_w=300)
done(s)

# =============================================================== P5
s = sol_slide("풀이 P5 (a)–(b) · 중심 영상전하", "중심 전하는 등전위를 깨지 않고 총전하를 맞춘다", "§3.2.2",
              "중심의 점전하는 구면 위 모든 점에서 같은 전위를 주므로 등전위 조건을 유지한다. 구면 전위 q/(4πε0 a)는 q가 중심에 만드는 전위와 같다.")
sphere_diagram(s, cx=220, cy=370, R=135, ax=430, image=True, center="q_{0}′", ground=False)
stack(s, [("t", "(a) 총전하 $0=q′+q_{0}′$", 22, True),
          ("e", "p5q0", r"q_0'=-q'=+\frac{R}{a}\,q\quad(\text{at }r=0)"),
          ("t", "(b) P4의 전위에 중심 전하의 전위를 더한다", 22, True),
          ("e", "p5V", r"V=V_{\rm P4}(r,\theta)+\frac{1}{4\pi\epsilon_0}\frac{qR/a}{r}"),
          ("t", "$r=R$에서 $V_{P4}=0$이므로 셋째 항만 남아 $θ$와 무관하다.", 22)],
      656, 150, 560)
result_box(s, "p5res", r"V(R,\theta)=\frac{q}{4\pi\epsilon_0 a}=\text{const.}", left=656, width=560)
done(s)

s = sol_slide("풀이 P5 (c) · 먼 거리 거동", "홀극 항의 계수가 전체 전하를 드러낸다", "§3.2.2, §3.4 예고",
              "가장 낮은 차수 항은 홀극 항 Q_total/(4πε0 r). 다중극 전개는 §3.4에서 다루므로 여기서는 예고로 언급한다.")
equation(s, "p5mono", r"V\simeq\frac{1}{4\pi\epsilon_0}\frac{Q_{\rm total}}{r}\qquad(r\to\infty)", 64, 150, 1152, align="center")
cols = [("접지된 구 (P4)", "p5g", r"Q_{\rm total}=q+q'=q\left(1-\frac{R}{a}\right)", "접지선으로 −qR/a가 들어온다"),
        ("중성 고립구 (P5)", "p5n", r"Q_{\rm total}=q+q'+q_0'=q", "도체 알짜 전하 0 → 멀리서는 q만 보인다")]
for i, (t1, k, l, t2) in enumerate(cols):
    x = 64 + i * 590
    text(s, t1, x, 260, 560, 36, size=24, bold=True, color="blue" if i else "navy")
    hline(s, x, 304, 562, "blue" if i else "navy", 3)
    equation(s, k, l, x, 330, 562)
    text(s, t2, x, 440, 562, 34, size=21, color="graydark")
result_box(s, "p5res2", r"\frac{q(1-R/a)}{4\pi\epsilon_0\,r}\ \longrightarrow\ \frac{q}{4\pi\epsilon_0\,r}", tail="둘 다 $1/r$ 차수, 계수가 $q(1−R/a)$에서 $q$로 커진다", tail_w=560)
done(s)

# =============================================================== P6
s = sol_slide("풀이 P6 (a) · 변수 분리", "곱 형태를 대입하면 두 상미분방정식으로 나뉜다", "§3.3.1",
              "y 방향 경계조건 Y(0)=Y(a)=0을 만족하려면 Y가 진동해야 하므로 Y 쪽 분리상수를 −k^2로 둔다.")
y = stack(s, [("e", "p6L", r"\frac{\partial^2V}{\partial x^2}+\frac{\partial^2V}{\partial y^2}=0,\qquad V=X(x)Y(y)", SCALE, "center")], 64, 150, 1152)
step_label(s, 1, 64, y + 14)
y = stack(s, [("t", "$XY$로 나누면 각 항은 서로 다른 변수만의 함수이므로 상수여야 한다", 22),
              ("e", "p6s", r"\frac{1}{X}\frac{\mathrm d^2X}{\mathrm dx^2}+\frac{1}{Y}\frac{\mathrm d^2Y}{\mathrm dy^2}=0")], 108, y + 14, 1100)
step_label(s, 2, 64, y + 6)
stack(s, [("t", "$y$ 방향의 두 영 경계가 진동해를 요구하므로 $Y$ 쪽에 $−k^{2}$", 22)], 108, y + 6, 1100)
result_box(s, "p6res", r"\frac{\mathrm d^2X}{\mathrm dx^2}=k^2X,\qquad \frac{\mathrm d^2Y}{\mathrm dy^2}=-k^2Y")
done(s)

s = sol_slide("풀이 P6 (b) · 경계조건으로 해 고르기", "경계가 한 모드뿐이면 해도 한 항이다", "§3.3.1",
              "Y(0)=Y(a)=0 → sin(nπy/a). X(0)=0 → sinh(kx). x=b 경계가 n=1 모드뿐이므로 계수 하나만 남는다.")
rect_diagram(s, x0=120, y0=240, w=260, h=250, right="V_{0} sin(πy/a)")
stack(s, [("t", "① $Y(0)=Y(a)=0$ → $Y=sin(nπy/a)$, $k=nπ/a$", 22),
          ("t", "② $X(0)=0$ → $X=sinh(kx)$", 22),
          ("t", "③ $x=b$의 경계는 $n=1$ 모드 하나", 22),
          ("e", "p6m", r"A_1\sinh\frac{\pi b}{a}=V_0")],
      656, 160, 560, gap=18)
result_box(s, "p6res2", r"V(x,y)=V_0\,\frac{\sinh(\pi x/a)}{\sinh(\pi b/a)}\,\sin\frac{\pi y}{a}", left=656, width=560)
done(s)

s = sol_slide("풀이 P6 (c) · 중심의 전위와 전기장", "중심의 전기장은 −x 방향만 남는다", "§3.3.1",
              "sinh 2u = 2 sinh u cosh u (u = πb/2a). y=a/2에서 ∂V/∂y ∝ cos(π/2)=0. 0 < V(중심) < V0/2로 최대·최소 원리와 부합.")
y = stack(s, [("e", "p6c1", r"V\!\left(\frac b2,\frac a2\right)=V_0\frac{\sinh(\pi b/2a)}{\sinh(\pi b/a)}=\frac{V_0}{2\cosh(\pi b/2a)}", SCALE, "center")], 64, 150, 1152)
step_label(s, 1, 64, y + 18)
y = stack(s, [("e", "p6c2", r"\frac{\partial V}{\partial x}=\frac{\pi V_0}{a}\frac{\cosh(\pi x/a)}{\sinh(\pi b/a)}\sin\frac{\pi y}{a}\ \to\ \frac{\pi V_0}{2a\,\sinh(\pi b/2a)}")], 108, y + 18, 1100)
step_label(s, 2, 64, y + 8)
stack(s, [("t", "$∂V/∂y ∝ cos(π/2)=0$ → $y$ 성분 없음. 전위가 높은 $x=b$에서 멀어지는 방향", 22)], 108, y + 8, 1100)
result_box(s, "p6res3", r"\mathbf E\!\left(\frac b2,\frac a2\right)=-\frac{\pi V_0}{2a\,\sinh(\pi b/2a)}\,\hat{\mathbf x}")
done(s)

# =============================================================== P7
s = sol_slide("풀이 P7 (a) · 푸리에 사인 계수", "짝수 모드는 cos nπ = 1 때문에 사라진다", "§3.3.1, Example 3.3",
              "직교성 ∫ sin(nπy/a) sin(n′πy/a) dy = (a/2)δ_nn′ 로 계수를 구한다. 상수 함수는 y=a/2에 대해 대칭이고 짝수 모드는 반대칭.")
y = stack(s, [("e", "p7e", r"V_0=\sum_{n=1}^{\infty}c_n\sin\frac{n\pi y}{a}", SCALE, "center")], 64, 150, 1152)
step_label(s, 1, 64, y + 16)
y = stack(s, [("t", "양변에 $sin(n′πy/a)$를 곱해 $0$부터 $a$까지 적분 (직교성)", 22),
              ("e", "p7c", r"c_n=\frac{2}{a}\int_0^a V_0\sin\frac{n\pi y}{a}\,\mathrm dy=\frac{2V_0}{n\pi}\,(1-\cos n\pi)")], 108, y + 16, 1100)
step_label(s, 2, 64, y + 8)
stack(s, [("t", "짝수 $n$: $cos nπ=1$ → $c_{n}=0$ · 홀수 $n$: $cos nπ=−1$", 22)], 108, y + 8, 1100)
result_box(s, "p7res", r"c_n=\begin{cases}\dfrac{4V_0}{n\pi}, & n=1,3,5,\dots\\[4pt] 0, & n=2,4,6,\dots\end{cases}", top=548, height=112)
done(s)

s = sol_slide("풀이 P7 (b)–(c) · 급수해와 확인", "각 모드를 경계 계수로 겹친다", "§3.3.1",
              "라플라스 방정식은 선형이므로 각 항이 만족하면 합도 만족한다. x=b에서는 (a)의 급수가 되어 V0를 재현한다.")
y = stack(s, [("e", "p7V", r"V(x,y)=\frac{4V_0}{\pi}\sum_{n=1,3,5,\dots}\frac{1}{n}\,\frac{\sinh(n\pi x/a)}{\sinh(n\pi b/a)}\,\sin\frac{n\pi y}{a}", SCALE, "center")], 64, 150, 1152)
hline(s, 64, y + 4, 1152, "hair", 1)
checks = [("라플라스 방정식", "p7k1", r"\nabla^2\!\left[\sinh(k_nx)\sin(k_ny)\right]=(k_n^2-k_n^2)\,(\cdots)=0"),
          ("세 영 경계", "p7k2", r"x=0:\ \sinh0=0,\quad y=0:\ \sin0=0,\quad y=a:\ \sin n\pi=0")]
yy = y + 24
for i, (t1, k, l) in enumerate(checks):
    step_label(s, i + 1, 64, yy)
    text(s, t1, 108, yy, 300, 34, size=22, bold=True)
    h = equation(s, k, l, 108, yy + 42, 1100)
    yy += h + 64
result_box(s, "p7res2", r"\nabla^2V=0,\qquad V(0,y)=V(x,0)=V(x,a)=0", tail="← 모든 항이 만족 (선형 중첩)", tail_w=380)
done(s)

# 수렴 그림
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
plt.rcParams["mathtext.fontset"] = "cm"
yv = np.linspace(0, 1, 1200)
fig, ax = plt.subplots(figsize=(6.2, 3.9), dpi=220)
for N, col, lw in ((1, "#98A2B3", 1.6), (5, "#344054", 1.6), (41, "#2563EB", 2.0)):
    f = sum(4 / (n * np.pi) * np.sin(n * np.pi * yv) for n in range(1, N + 1, 2))
    ax.plot(yv, f, color=col, lw=lw, label=f"$n\\leq{N}$")
ax.axhline(1, color="#191919", lw=1, ls="--")
ax.set_xlabel("$y/a$", fontsize=18); ax.set_ylabel("$V(b,y)/V_0$", fontsize=18)
ax.set_xlim(0, 1); ax.set_ylim(-0.05, 1.3)
ax.tick_params(labelsize=15)
for sp in ("top", "right"): ax.spines[sp].set_visible(False)
ax.legend(frameon=False, fontsize=16, loc="lower center", ncol=3)
fig.tight_layout(); fig.savefig(EQ / "p7_partial_sums.png", transparent=True); plt.close(fig)

s = sol_slide("풀이 P7 (d) · 불연속점에서의 수렴", "불연속점에서는 양쪽 극한의 평균으로 수렴한다", "§3.3.1",
              "사인 급수는 [−a, a]로의 홀함수 연장의 푸리에 급수. y=0, a에서 (V0 + (−V0))/2 = 0. 불연속점 근처의 약 9% 넘침은 깁스 현상. 물리적으로는 모서리에 가는 절연 간격이 있다고 본다.")
s.shapes.add_picture(str(EQ / "p7_partial_sums.png"), px(64), px(150), px(560))
text(s, "그림: $x=b$에서 부분합 (홀수 모드 $n≤N$)", 64, 512, 560, 30, size=17, color="gray", align="center")
stack(s, [("t", "모서리 $(b,0)$, $(b,a)$에서 경계값이 $V_{0}$에서 $0$으로 끊어진다.", 22),
          ("t", "푸리에 급수는 불연속점에서 양쪽 극한의 평균으로 수렴한다. 사인 급수는 홀함수 연장의 급수이므로 $y=0, a$에서는", 22),
          ("e", "p7avg", r"\tfrac12\left[V_0+(-V_0)\right]=0"),
          ("t", "불연속점 근처에서는 약 9% 넘침(깁스 현상). 영역 안($x<b$)에서는 $sinh$ 인자 때문에 빠르게 수렴한다.", 22)],
      656, 145, 560, gap=8)
result_box(s, "p7res3", r"\frac12\left[f(y^-)+f(y^+)\right]", tail="불연속점에서 급수가 수렴하는 값 (모서리 $y=0, a$에서는 $0$)", tail_w=650)
done(s)

# =============================================================== P8
s = sol_slide("풀이 P8 (a)–(b) · 안쪽 해와 바깥쪽 해", "경계값 V₀ cos θ는 ℓ = 1 항 하나만 고른다", "§3.3.2",
              "르장드르 다항식의 직교성으로 ℓ=1 계수만 0이 아니다. 안쪽은 원점 유한 조건으로 B_ℓ=0, 바깥은 무한대 조건으로 A_ℓ=0.")
cos_sphere(s, cx=250, cy=400, R=140, inside_field=True)
stack(s, [("e", "p8gen", r"V=\sum_{\ell}\left(A_\ell r^\ell+\frac{B_\ell}{r^{\ell+1}}\right)P_\ell(\cos\theta)", 2.3),
          ("t", "(a) $r<R$: 원점 유한 → $B_{ℓ}=0$, $A_{1}R=V_{0}$", 22, True),
          ("e", "p8in", r"V=V_0\frac{r}{R}\cos\theta=\frac{V_0}{R}\,z,\qquad \mathbf E=-\frac{V_0}{R}\hat{\mathbf z}"),
          ("t", "(b) $r>R$: $V→0$ → $A_{ℓ}=0$, $B_{1}/R^{2}=V_{0}$", 22, True),
          ("e", "p8out", r"V=V_0\frac{R^2}{r^2}\cos\theta")],
      600, 150, 616, gap=12)
result_box(s, "p8res", r"\mathbf E_{\rm in}=-\frac{V_0}{R}\hat{\mathbf z},\qquad V_{\rm out}=V_0\frac{R^2}{r^2}\cos\theta", left=600, width=616)
done(s)

s = sol_slide("풀이 P8 (c)–(d) · 모드와 쌍극자 모멘트", "바깥 해는 정확히 쌍극자 전위이다", "§3.3.2, §3.4.3 예고",
              "바깥 장은 쌍극자 장. 구면의 면전하는 σ = ε0(E_r^out − E_r^in) = 3ε0V0 cosθ / R 로 검산할 수 있다.")
y = stack(s, [("t", "(c) 두 해 모두 $ℓ=1$, $P_{1}(cos θ)=cos θ$ 하나만 필요하다.", 23, True),
              ("g", 8),
              ("t", "(d) 쌍극자 전위와 비교", 23, True),
              ("e", "p8dip", r"V_{\rm dip}=\frac{1}{4\pi\epsilon_0}\frac{p\cos\theta}{r^2}\quad\Longleftrightarrow\quad V_0\frac{R^2}{r^2}\cos\theta", SCALE, "center"),
              ("g", 6),
              ("t", "검산", 21, True),
              ("e", "p8chk", r"r=R:\ V_{\rm in}=V_{\rm out}=V_0\cos\theta,\qquad \sigma=\epsilon_0\left(E_r^{\rm out}-E_r^{\rm in}\right)\Big|_R=\frac{3\epsilon_0V_0}{R}\cos\theta", 2.25, "center")],
          64, 150, 1152)
result_box(s, "p8res2", r"\ell=1,\qquad \mathbf p=4\pi\epsilon_0R^2V_0\,\hat{\mathbf z}")
done(s)

# =============================================================== summary
s = new_slide(); header(s, "W07 / FINAL CHECK", "최종 검산 체크리스트")
cols = [("영상법 (P1–P5)", ["실제 영역을 먼저 적었다", "영상전하가 실제 영역 밖에 있다", "도체 표면에서 경계 전위를 만족한다", "접지와 고립 도체의 조건을 구분했다", "에너지에 계수 1/2를 적용했다"]),
        ("변수 분리법 (P6–P8)", ["분리상수의 부호를 경계조건으로 정했다", "각 항이 라플라스 방정식을 만족한다", "영 경계를 모든 항이 만족한다", "남은 경계를 직교성으로 맞췄다", "불연속점 수렴값과 먼 거리 거동을 확인했다"])]
for i, (t1, its) in enumerate(cols):
    x = 64 + i * 590
    text(s, t1, x, 160, 560, 36, size=25, bold=True, color="blue")
    hline(s, x, 206, 562, "blue", 3)
    yy = 228
    for it in its:
        rect(s, x + 2, yy + 6, 20, 20, line="graydark", lw=1.5)
        text(s, it, x + 38, yy, 530, 34, size=22)
        yy += 70
notes(s, "연습 마무리. 영상법과 변수 분리법의 공통 점검 항목을 정리한다. 다음 주 다중극 전개(§3.4)로 연결.", "§3.1–§3.3")

prs.save(OUT)
print("slides:", len(prs.slides), "->", OUT)
